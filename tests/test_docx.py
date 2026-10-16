import asyncio
import sys
from pathlib import Path

# Añadir backend al path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "backend"))

import docx
from app.services.translators.mock_translator import MockTranslator
from app.services.parsers.docx_parser import DocxDocumentParser

async def test_docx_translation():
    print("--- 1. Creando documento Word de prueba ---")
    doc = docx.Document()
    doc.add_heading("Manual de Operaciones de Software", level=1)
    
    p = doc.add_paragraph("Este es un párrafo normal con ")
    p.add_run("texto en negrita").bold = True
    p.add_run(" y ")
    p.add_run("palabras en cursiva").italic = True
    p.add_run(".")
    
    # Agregar una tabla
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Parámetro"
    table.rows[0].cells[1].text = "Descripción"
    table.rows[1].cells[0].text = "Timeout"
    table.rows[1].cells[1].text = "Tiempo máximo de espera del servidor"

    test_input = Path("storage/uploads/documento_prueba.docx")
    test_output = Path("storage/processed/documento_prueba_traducido.docx")
    
    test_input.parent.mkdir(parents=True, exist_ok=True)
    test_output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(test_input))
    print(f"Documento de prueba guardado en: {test_input}")

    print("\n--- 2. Ejecutando Traducción (Modo Mock / Desarrollo) ---")
    translator = MockTranslator()
    parser = DocxDocumentParser(translator=translator)

    async def progress_logger(current, total, msg):
        print(f"Progreso [{current}/{total}]: {msg}")

    await parser.translate_document(
        input_path=test_input,
        output_path=test_output,
        source_lang="es",
        target_lang="en",
        progress_callback=progress_logger
    )

    print(f"\n--- 3. Verificando documento traducido en: {test_output} ---")
    doc_traducido = docx.Document(str(test_output))
    for i, p in enumerate(doc_traducido.paragraphs):
        print(f"Párrafo {i+1}: {p.text}")

    print("\n¡Prueba de traducción DOCX completada con ÉXITO!")

if __name__ == "__main__":
    asyncio.run(test_docx_translation())
