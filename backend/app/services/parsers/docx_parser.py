import re
import logging
from pathlib import Path
from typing import Callable, Awaitable, Optional, List
import docx
from docx.text.paragraph import Paragraph
from app.services.parsers.base import BaseDocumentParser

logger = logging.getLogger(__name__)

class DocxDocumentParser(BaseDocumentParser):
    """
    Parser para documentos Microsoft Word (.docx).
    Preserva estilos, tablas, negritas, cursivas y estructura del documento original.
    """

    async def translate_document(
        self,
        input_path: Path,
        output_path: Path,
        source_lang: str,
        target_lang: str,
        tone: Optional[str] = "natural",
        progress_callback: Optional[Callable[[int, int, str], Awaitable[None]]] = None
    ) -> Path:
        doc = docx.Document(str(input_path))

        # 1. Recolectar todos los párrafos (cuerpo principal, tablas y encabezados)
        paragraphs_to_process: List[Paragraph] = []

        # Párrafos del cuerpo principal
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs_to_process.append(p)

        # Párrafos dentro de tablas
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.text.strip():
                            paragraphs_to_process.append(p)

        # Párrafos de encabezados y pies de página
        for section in doc.sections:
            for p in section.header.paragraphs:
                if p.text.strip():
                    paragraphs_to_process.append(p)
            for p in section.footer.paragraphs:
                if p.text.strip():
                    paragraphs_to_process.append(p)

        total_paragraphs = len(paragraphs_to_process)
        if total_paragraphs == 0:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(output_path))
            return output_path

        if progress_callback:
            await progress_callback(0, total_paragraphs, "Iniciando análisis del documento...")

        # 2. Procesamiento por lotes (Batching) para optimizar velocidad y contexto
        BATCH_SIZE = 15
        processed_count = 0

        for i in range(0, total_paragraphs, BATCH_SIZE):
            batch = paragraphs_to_process[i : i + BATCH_SIZE]
            
            # Preparar textos con enmascarado de formato (runs)
            masked_texts = [self._mask_paragraph(p) for p in batch]
            
            # Traducir lote con tono especificado
            translated_masked_texts = await self.translator.translate_batch(
                texts=masked_texts,
                source_lang=source_lang,
                target_lang=target_lang,
                tone=tone
            )

            # Reinyectar texto traducido en cada párrafo
            for p, trans_text in zip(batch, translated_masked_texts):
                self._unmask_and_apply(p, trans_text)
                processed_count += 1

            if progress_callback:
                await progress_callback(
                    processed_count,
                    total_paragraphs,
                    f"Traducidos {processed_count} de {total_paragraphs} párrafos..."
                )

        # 3. Guardar el archivo final traducido
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        
        if progress_callback:
            await progress_callback(total_paragraphs, total_paragraphs, "Documento completado con éxito.")

        return output_path

    def _mask_paragraph(self, p: Paragraph) -> str:
        """
        Si un párrafo tiene múltiples runs con estilos diferentes, los envuelve en tags <r0>, <r1>.
        Si es un párrafo de estilo uniforme, devuelve el texto simple.
        """
        runs = [r for r in p.runs if r.text]
        if len(runs) <= 1:
            return p.text

        masked_parts = []
        for idx, run in enumerate(runs):
            if run.text:
                masked_parts.append(f"<r{idx}>{run.text}</r{idx}>")
        return "".join(masked_parts)

    def _unmask_and_apply(self, p: Paragraph, translated_text: str):
        """
        Aplica el texto traducido de vuelta al párrafo preservando el formato de los runs.
        """
        runs = [r for r in p.runs if r.text]
        if not runs:
            p.text = translated_text
            return

        tag_pattern = re.compile(r"<r(\d+)>(.*?)</r\1>", re.DOTALL)
        matches = list(tag_pattern.finditer(translated_text))

        if matches:
            run_translations = {}
            for m in matches:
                run_idx = int(m.group(1))
                run_content = m.group(2)
                run_translations[run_idx] = run_content

            for idx, run in enumerate(runs):
                if idx in run_translations:
                    run.text = run_translations[idx]
                else:
                    run.text = ""
        else:
            # Si no hubo tags en la respuesta, colocar en el primer run y vaciar los demás
            runs[0].text = translated_text
            for run in runs[1:]:
                run.text = ""
